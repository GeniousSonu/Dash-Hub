import csv
from datetime import datetime, timedelta, timezone
from pathlib import Path
import sqlite3
from typing import Any
import openpyxl

from backend.app.core.database import new_uuid, utc_now_iso
from backend.app.models.enums import BuilderType, HealthStatus, WebsiteEnvironment, WorkflowStatus
from backend.app.services.audit_service import AuditService


class InventoryImportService:
    @staticmethod
    def parse_file(file_path: Path | str) -> list[dict[str, Any]]:
        path = Path(file_path)
        rows: list[dict[str, Any]] = []
        if path.suffix.lower() == ".xlsx":
            wb = openpyxl.load_workbook(str(path), data_only=True)
            ws = wb.active
            headers = [str(cell.value or "").strip().lower().replace(" ", "_") for cell in next(ws.iter_rows())]
            for r in ws.iter_rows(min_row=2, values_only=True):
                if not any(r):
                    continue
                row_dict = {headers[i]: r[i] for i in range(min(len(headers), len(r))) if headers[i]}
                rows.append(row_dict)
        else:
            with open(path, "r", encoding="utf-8-sig") as f:
                reader = csv.DictReader(f)
                for r in reader:
                    cleaned = {k.strip().lower().replace(" ", "_"): (v or "").strip() for k, v in r.items() if k}
                    rows.append(cleaned)
        return rows

    @classmethod
    def import_websites(
        cls,
        conn: sqlite3.Connection,
        *,
        rows: list[dict[str, Any]],
        default_project_id: str,
        default_maintainer_id: str,
        actor_user_id: str,
        dry_run: bool = False,
    ) -> dict[str, Any]:
        now = utc_now_iso()
        imported = 0
        duplicates_skipped = 0
        created_websites = []

        existing_pairs = {
            (r["domain"].strip().lower(), r["environment"].strip().lower())
            for r in conn.execute("SELECT domain, environment FROM websites").fetchall()
        }

        for row in rows:
            name = (row.get("name") or row.get("site_name") or row.get("website_name") or "").strip()
            domain = (row.get("domain") or row.get("url") or "").strip().lower()
            domain = domain.replace("https://", "").replace("http://", "").split("/")[0]
            if not domain:
                continue
            if not name:
                name = domain

            env = (row.get("environment") or row.get("env") or "live").strip().lower()
            if env not in ("live", "staging", "dev"):
                env = "staging" if "staging" in domain else "live"

            pair = (domain, env)
            if pair in existing_pairs:
                duplicates_skipped += 1
                continue

            existing_pairs.add(pair)
            imported += 1
            site_id = new_uuid()
            url = f"https://{domain}"
            builder = (row.get("builder") or "gutenberg").strip().lower()
            if builder not in ("gutenberg", "divi", "elementor"):
                builder = "divi" if "divi" in builder else "gutenberg"

            hosting_type = (row.get("hosting_type") or "managed_wp").strip()
            frequency = (row.get("frequency") or "monthly").strip().lower()
            wp_version = str(row.get("wp_version") or "not_verified").strip()
            php_version = str(row.get("php_version") or "not_verified").strip()
            maintainer_id = row.get("maintainer_id") or default_maintainer_id

            next_due_at = (datetime.now(timezone.utc) + timedelta(days=30)).date().isoformat()

            site_data = {
                "id": site_id,
                "project_id": default_project_id,
                "name": name,
                "domain": domain,
                "url": url,
                "environment": env,
                "health_status": HealthStatus.UNKNOWN.value,
                "workflow_status": WorkflowStatus.DUE.value,
                "maintainer_id": maintainer_id,
                "frequency": frequency,
                "next_due_at": next_due_at,
                "builder": builder,
                "hosting_type": hosting_type,
                "wp_version": wp_version,
                "php_version": php_version,
            }
            created_websites.append(site_data)

            if not dry_run:
                conn.execute(
                    """
                    INSERT INTO websites (
                        id, project_id, name, domain, url, environment, health_status,
                        workflow_status, maintainer_id, frequency, next_due_at,
                        wp_version, php_version, builder, hosting_type, created_at, updated_at
                    ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                    """,
                    (
                        site_id,
                        default_project_id,
                        name,
                        domain,
                        url,
                        env,
                        HealthStatus.UNKNOWN.value,
                        WorkflowStatus.DUE.value,
                        maintainer_id,
                        frequency,
                        next_due_at,
                        wp_version,
                        php_version,
                        builder,
                        hosting_type,
                        now,
                        now,
                    ),
                )
                conn.execute(
                    """
                    INSERT INTO maintenance_schedules (
                        id, website_id, template_code, frequency, interval_days, anchor_policy,
                        planned_due_at, next_due_at, owner_id, is_active, updated_at
                    ) VALUES (?, ?, 'WP_MAINT_SOP', ?, 30, 'planned_date', ?, ?, ?, 1, ?)
                    """,
                    (
                        new_uuid(),
                        site_id,
                        frequency,
                        next_due_at,
                        next_due_at,
                        maintainer_id,
                        now,
                    ),
                )
                AuditService.log_event(
                    conn,
                    actor_user_id=actor_user_id,
                    entity_type="website",
                    entity_id=site_id,
                    action="website_imported",
                    summary=f"Imported website {name} ({domain}, {env})",
                    after_state=site_data,
                )

        return {
            "imported": imported,
            "duplicates_skipped": duplicates_skipped,
            "total_rows": len(rows),
            "created_websites": created_websites,
            "dry_run": dry_run,
        }
